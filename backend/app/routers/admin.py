import io
import re
from datetime import datetime
from typing import List, Optional
from fastapi import APIRouter, Depends, HTTPException, Query, status
from fastapi.responses import Response
from openpyxl import Workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter
from sqlalchemy import func
from sqlalchemy.orm import Session
from .. import schemas
from .. import database
from .. import models
from .. import security
from .adopciones import ESTADO_APROBADA

router = APIRouter(prefix="/admin", tags=["Administración"], dependencies=[Depends(security.require_admin)])

MESES = ["Enero", "Febrero", "Marzo", "Abril", "Mayo", "Junio", "Julio", "Agosto",
         "Septiembre", "Octubre", "Noviembre", "Diciembre"]
PATRON_MES = r"^\d{4}-(0[1-9]|1[0-2])$"


def _rango_mes(mes: Optional[str]):
    """'2026-10' -> (1 oct 2026, 1 nov 2026, '2026-10'). Sin dato = mes actual."""
    if mes is None:
        hoy = datetime.now()
        anio, num = hoy.year, hoy.month
    else:
        anio, num = int(mes[:4]), int(mes[5:7])
    inicio = datetime(anio, num, 1)
    fin = datetime(anio + 1, 1, 1) if num == 12 else datetime(anio, num + 1, 1)
    return inicio, fin, f"{anio}-{num:02d}"


def _calcular_metricas(db: Session, mes: Optional[str]) -> schemas.MetricasOut:
    inicio, fin, etiqueta = _rango_mes(mes)
    S, O, U, M = models.SolicitudAdopcion, models.OrdenCompra, models.Usuario, models.Mascota

    def contar(consulta):
        return consulta.scalar() or 0

    en_mes_sol = (S.fecha_solicitud >= inicio, S.fecha_solicitud < fin)
    en_mes_ord = (O.fecha_compra >= inicio, O.fecha_compra < fin)
    return schemas.MetricasOut(
        mes=etiqueta,
        solicitudes_mes=contar(db.query(func.count(S.id)).filter(*en_mes_sol)),
        adopciones_mes=contar(db.query(func.count(S.id)).filter(S.estado == ESTADO_APROBADA, *en_mes_sol)),
        pedidos_mes=contar(db.query(func.count(O.id)).filter(*en_mes_ord)),
        ventas_mes=float(db.query(func.coalesce(func.sum(O.total), 0)).filter(*en_mes_ord).scalar() or 0),
        ventas_totales=float(db.query(func.coalesce(func.sum(O.total), 0)).scalar() or 0),
        usuarios_registrados=contar(db.query(func.count(U.id))),
        usuarios_nuevos_mes=contar(db.query(func.count(U.id)).filter(U.fecha_registro >= inicio, U.fecha_registro < fin)),
        mascotas_disponibles=contar(db.query(func.count(M.id)).filter(M.estado_adopcion == "Disponible")),
        mascotas_adoptadas=contar(db.query(func.count(M.id)).filter(M.estado_adopcion == "Adoptada")),
    )


@router.get("/metricas", response_model=schemas.MetricasOut)
def metricas(mes: Optional[str] = Query(None, pattern=PATRON_MES), db: Session = Depends(database.get_db)):
    return _calcular_metricas(db, mes)


# ---------- Reporte en Excel (con formato) ----------
VERDE = "28A745"
BORDE = Border(*(Side(style="thin", color="D0D5DA"),) * 4)


def _encabezado(hoja, fila, titulos):
    for col, titulo in enumerate(titulos, start=1):
        c = hoja.cell(row=fila, column=col, value=titulo)
        c.font = Font(bold=True, color="FFFFFF")
        c.fill = PatternFill("solid", fgColor=VERDE)
        c.alignment = Alignment(horizontal="center", vertical="center")
        c.border = BORDE


def _ajustar_columnas(hoja, minimo=12, maximo=45):
    for col in hoja.columns:
        largo = max((len(str(c.value)) for c in col if c.value is not None and c.row > 4), default=minimo)
        hoja.column_dimensions[get_column_letter(col[0].column)].width = max(minimo, min(maximo, largo + 3))


@router.get("/reporte.xlsx")
def reporte_excel(mes: Optional[str] = Query(None, pattern=PATRON_MES), db: Session = Depends(database.get_db)):
    inicio, fin, etiqueta = _rango_mes(mes)
    m = _calcular_metricas(db, etiqueta)
    nombre_mes = f"{MESES[inicio.month - 1]} {inicio.year}"
    dinero = '"$"#,##0.00'
    fecha_fmt = "dd/mm/yyyy hh:mm"

    wb = Workbook()

    # ---- Hoja 1: Resumen
    h = wb.active
    h.title = "Resumen"
    h.sheet_view.showGridLines = False
    h["A1"] = "Paws&Shop · Reporte mensual"
    h["A1"].font = Font(size=18, bold=True, color=VERDE)
    h["A2"] = f"Periodo: {nombre_mes}"
    h["A2"].font = Font(size=12, bold=True)
    h["A3"] = f"Generado el {datetime.now():%d/%m/%Y a las %H:%M}"
    h["A3"].font = Font(italic=True, color="666666")
    _encabezado(h, 5, ["Indicador", "Valor"])
    filas = [
        (f"Solicitudes de adopción ({nombre_mes})", m.solicitudes_mes, None),
        (f"Adopciones aprobadas ({nombre_mes})", m.adopciones_mes, None),
        (f"Pedidos de la tienda ({nombre_mes})", m.pedidos_mes, None),
        (f"Ventas de {nombre_mes}", m.ventas_mes, dinero),
        ("Ventas acumuladas (todos los meses)", m.ventas_totales, dinero),
        (f"Usuarios nuevos ({nombre_mes})", m.usuarios_nuevos_mes, None),
        ("Usuarios registrados (total)", m.usuarios_registrados, None),
        ("Mascotas disponibles hoy", m.mascotas_disponibles, None),
        ("Mascotas adoptadas (total)", m.mascotas_adoptadas, None),
    ]
    for i, (texto, valor, fmt) in enumerate(filas, start=6):
        a, b = h.cell(row=i, column=1, value=texto), h.cell(row=i, column=2, value=valor)
        a.border = b.border = BORDE
        b.alignment = Alignment(horizontal="right")
        if fmt:
            b.number_format = fmt
        if i % 2 == 1:
            a.fill = b.fill = PatternFill("solid", fgColor="F3F8F4")
    h.column_dimensions["A"].width = 48
    h.column_dimensions["B"].width = 20

    # ---- Hoja 2: Adopciones del mes
    h = wb.create_sheet("Adopciones")
    h["A1"] = f"Solicitudes de adopción · {nombre_mes}"
    h["A1"].font = Font(size=14, bold=True, color=VERDE)
    _encabezado(h, 3, ["Fecha", "Adoptante", "Mascota", "Refugio", "Estado", "Teléfono", "Vivienda"])
    S, M, U, R = models.SolicitudAdopcion, models.Mascota, models.Usuario, models.Refugio
    filas = (db.query(S, M, U, R)
             .join(M, M.id == S.mascota_id).join(U, U.id == S.usuario_id)
             .outerjoin(R, R.id == M.refugio_id)
             .filter(S.fecha_solicitud >= inicio, S.fecha_solicitud < fin)
             .order_by(S.fecha_solicitud).all())
    for i, (s, mas, usu, ref) in enumerate(filas, start=4):
        valores = [s.fecha_solicitud, usu.nombre_completo, mas.nombre, ref.nombre if ref else "", s.estado,
                   s.telefono or "", s.tipo_vivienda or ""]
        for col, v in enumerate(valores, start=1):
            c = h.cell(row=i, column=col, value=v)
            c.border = BORDE
            if col == 1:
                c.number_format = fecha_fmt
    if not filas:
        h["A4"] = "Sin solicitudes en este periodo."
        h["A4"].font = Font(italic=True, color="666666")
    h.freeze_panes = "A4"
    if filas:
        h.auto_filter.ref = f"A3:G{3 + len(filas)}"
    _ajustar_columnas(h)

    # ---- Hoja 3: Ventas del mes
    h = wb.create_sheet("Ventas")
    h["A1"] = f"Ventas de la tienda · {nombre_mes}"
    h["A1"].font = Font(size=14, bold=True, color=VERDE)
    _encabezado(h, 3, ["N.º pedido", "Fecha", "Cliente", "Estado", "Total"])
    O = models.OrdenCompra
    ordenes = (db.query(O, U).join(U, U.id == O.usuario_id)
               .filter(O.fecha_compra >= inicio, O.fecha_compra < fin).order_by(O.fecha_compra).all())
    for i, (o, usu) in enumerate(ordenes, start=4):
        for col, v in enumerate([o.id, o.fecha_compra, usu.nombre_completo, o.estado_pedido, float(o.total)], start=1):
            c = h.cell(row=i, column=col, value=v)
            c.border = BORDE
            if col == 2:
                c.number_format = fecha_fmt
            if col == 5:
                c.number_format = dinero
    if ordenes:
        ultima = 3 + len(ordenes)
        h.cell(row=ultima + 1, column=4, value="TOTAL").font = Font(bold=True)
        t = h.cell(row=ultima + 1, column=5, value=f"=SUM(E4:E{ultima})")
        t.font = Font(bold=True)
        t.number_format = dinero
        h.auto_filter.ref = f"A3:E{ultima}"
    else:
        h["A4"] = "Sin ventas en este periodo."
        h["A4"].font = Font(italic=True, color="666666")
    h.freeze_panes = "A4"
    _ajustar_columnas(h)

    buffer = io.BytesIO()
    wb.save(buffer)
    return Response(
        content=buffer.getvalue(),
        media_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
        headers={"Content-Disposition": f"attachment; filename=reporte_pawsshop_{etiqueta}.xlsx"},
    )


@router.get("/auditoria", response_model=List[schemas.AuditoriaOut])
def auditoria(db: Session = Depends(database.get_db)):
    return db.query(models.Auditoria).order_by(models.Auditoria.id.desc()).limit(200).all()


@router.get("/encargados", response_model=List[schemas.EncargadoOut])
def listar_encargados(db: Session = Depends(database.get_db)):
    filas = (db.query(models.Usuario, models.Refugio)
             .outerjoin(models.Refugio, models.Refugio.id == models.Usuario.refugio_id)
             .filter(models.Usuario.rol_id == security.ROL_ENCARGADO)
             .order_by(models.Usuario.id).all())
    return [schemas.EncargadoOut(id=u.id, nombre_completo=u.nombre_completo, email=u.email,
                                 refugio_id=u.refugio_id, refugio_nombre=r.nombre if r else None) for u, r in filas]


@router.post("/encargados", status_code=status.HTTP_201_CREATED, response_model=schemas.EncargadoOut)
def crear_encargado(datos: schemas.EncargadoCreate, db: Session = Depends(database.get_db)):
    if db.query(models.Usuario).filter(models.Usuario.email == datos.email).first():
        raise HTTPException(status_code=status.HTTP_400_BAD_REQUEST, detail="El correo ya está registrado")
    refugio = db.query(models.Refugio).filter(models.Refugio.id == datos.refugio_id).first()
    if not refugio:
        raise HTTPException(status_code=status.HTTP_404_NOT_FOUND, detail="El refugio no existe")

    nuevo = models.Usuario(
        nombre_completo=datos.nombre_completo, email=datos.email,
        password_hash=security.hash_password(datos.password),
        rol_id=security.ROL_ENCARGADO, refugio_id=refugio.id,
    )
    db.add(nuevo)
    db.commit()
    db.refresh(nuevo)
    return schemas.EncargadoOut(id=nuevo.id, nombre_completo=nuevo.nombre_completo, email=nuevo.email,
                                refugio_id=refugio.id, refugio_nombre=refugio.nombre)